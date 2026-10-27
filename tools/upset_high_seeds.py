import os
import csv
import glob
import pdem_ratings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# The higher seeds that were upset in the NCAA tournament, 2010-2026:   python tools/upset_high_seeds.py
# Every game the lower seed won by 4+ seed lines (reports/pdem_seed_gap_4plus.csv, First Four not included), from the
# beaten higher seed's side: its November record, its November record against that season's pre-tournament top 40 (by net
# rating, from the pre-tournament page) and the round it went out in (the upset itself). Rows are highlighted:
# yellow = 2+ November losses, orange = no November games against the top 40, red = both. Writes reports/upset_high_seeds.xlsx.

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS = os.path.join(ROOT, "reports")
PAGES = os.path.join(ROOT, "dashboard")
TOP = 40
ROUNDS = {"First round": "Round of 64", "Second round": "Round of 32", "Sweet 16": "Sweet 16", "Elite Eight": "Elite Eight",
          "Final Four": "Final Four", "Championship": "Championship"}

def seasons():
    """season -> (team -> its pre-tournament games, the pre-tournament top 40)"""
    out = {}
    for path in [os.path.join(PAGES, "dashboard.html")] + glob.glob(os.path.join(PAGES, "dashboard_*.html")):
        if path.endswith(("_full.html", "_template.html", "upsets.html")): continue
        data = pdem_ratings.page_data(path)
        rated = sorted((t for t in data["teams"] if t.get("n") is not None), key=lambda t: -t["n"])
        out[data["meta"]["season"]] = ({t["name"]: t.get("sched", []) for t in data["teams"]}, {t["name"] for t in rated[:TOP]})
    return out

def main():
    pages = seasons()
    rows = []
    for g in csv.DictReader(open(os.path.join(REPORTS, "pdem_seed_gap_4plus.csv"), encoding="utf-8")):
        if g["lower_won"] != "yes": continue
        sched, top = pages[g["season"]]
        games = sched.get(g["higher_team"], [])
        nov = [x for x in games if x["d"][5:7] == "11"]
        vs = [x for x in nov if x["opp"] in top]          # November games against the top 40 only
        rows.append([int(g["date"][:4]), ROUNDS[g["round"]], int(g["higher_seed"]), g["higher_team"], int(g["lower_seed"]), g["lower_team"],
                     f'{g["lower_pts"]}–{g["higher_pts"]}', sum(x["pf"] > x["pa"] for x in nov), sum(x["pf"] < x["pa"] for x in nov),
                     None, sum(x["pf"] > x["pa"] for x in vs), sum(x["pf"] < x["pa"] for x in vs)])     # None: the record formula's column
    order = list(ROUNDS.values())
    rows.sort(key=lambda r: (-order.index(r[1]), r[0], r[2]))      # deepest exit round first, then year, then seed

    wb = Workbook()
    ws = wb.active
    ws.title = "Upset higher seeds"
    font = lambda **k: Font(name="Arial", **k)
    yellow, orange, red = "FFF2A8", "FFC98B", "F4A6A6"
    ws["A1"] = "Higher seeds upset in the NCAA tournament by a team 4+ seed lines worse, 2010–2026"
    ws["A1"].font = font(bold=True, size=13)
    ws["A2"] = (f"{len(rows)} games (First Four not included). November = the team's games in November of that season. "
                f"Top 40 = that season's pre-tournament top {TOP} by Strongside net rating; the top-40 record counts November games only.")
    ws["A2"].font = font(size=9, italic=True)
    key = [("Key:", None), ("2+ November losses", yellow), ("No November games vs top 40 (not battle tested)", orange), ("Both", red)]
    for i, (text, fill) in enumerate(key):
        c = ws.cell(row=3, column=1 + i * 2, value=text)
        c.font = font(size=9, bold=fill is None)
        if fill: c.fill = PatternFill("solid", fgColor=fill)

    head = ["Year", "Exit round", "Seed", "Team (upset)", "Upset by seed", "Upset by", "Score", "Nov W", "Nov L", "November record",
            "Nov top-40 W", "Nov top-40 L", "November record vs top 40", "Nov top-40 games"]
    H = 5
    for j, h in enumerate(head, 1):
        c = ws.cell(row=H, column=j, value=h)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14213D")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D5D9E0")
    for i, r in enumerate(rows):
        n = H + 1 + i
        for j, v in enumerate(r, 1):
            if v is not None: ws.cell(row=n, column=j, value=v)
        ws.cell(row=n, column=10, value=f'=H{n}&"-"&I{n}')            # November record, e.g. 6-1
        ws.cell(row=n, column=13, value=f'=K{n}&"-"&L{n}')            # record vs top 40, e.g. 3-5
        ws.cell(row=n, column=14, value=f"=K{n}+L{n}")
        for j in range(1, len(head) + 1):
            c = ws.cell(row=n, column=j)
            c.font = font(size=10)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal="left" if j in (2, 4, 6) else "center")
    last = H + len(rows)
    rng = f"A{H + 1}:{get_column_letter(len(head))}{last}"
    # the row colours follow the numbers: red first, so a team meeting both rules is red
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND($I{H + 1}>=2,$N{H + 1}=0)"], fill=PatternFill("solid", fgColor=red, bgColor=red), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$I{H + 1}>=2"], fill=PatternFill("solid", fgColor=yellow, bgColor=yellow), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$N{H + 1}=0"], fill=PatternFill("solid", fgColor=orange, bgColor=orange), stopIfTrue=True))
    # totals under the table
    t = last + 2
    ws.cell(row=t, column=1, value="Totals").font = font(bold=True)
    ws.cell(row=t, column=8, value=f"=SUM(H{H + 1}:H{last})")
    ws.cell(row=t, column=9, value=f"=SUM(I{H + 1}:I{last})")
    ws.cell(row=t, column=10, value=f'=H{t}&"-"&I{t}')
    ws.cell(row=t, column=11, value=f"=SUM(K{H + 1}:K{last})")
    ws.cell(row=t, column=12, value=f"=SUM(L{H + 1}:L{last})")
    ws.cell(row=t, column=13, value=f'=K{t}&"-"&L{t}')
    ws.cell(row=t + 1, column=1, value="Teams with 2+ November losses")
    ws.cell(row=t + 1, column=9, value=f'=COUNTIF(I{H + 1}:I{last},">=2")')
    ws.cell(row=t + 2, column=1, value="Teams with no November games vs top 40")
    ws.cell(row=t + 2, column=14, value=f"=COUNTIF(N{H + 1}:N{last},0)")
    ws.cell(row=t + 3, column=1, value="Both")
    ws.cell(row=t + 3, column=14, value=f'=COUNTIFS(I{H + 1}:I{last},">=2",N{H + 1}:N{last},0)')
    for rr in range(t, t + 4):
        for j in range(1, len(head) + 1):
            ws.cell(row=rr, column=j).font = font(size=10, bold=rr == t)
    widths = [10.9, 16, 17.7, 19, 40.3, 16.7, 11.6, 9.7, 11.6, 15.3, 13.6, 13.6, 15.3, 13.6]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 16.5
    ws.row_dimensions[H].height = 48.75
    ws.freeze_panes = f"A{H + 1}"
    ws.auto_filter.ref = f"A{H}:{get_column_letter(len(head))}{last}"
    path = os.path.join(REPORTS, "upset_high_seeds.xlsx")
    wb.save(path)
    print(len(rows), "rows ->", path)

if __name__ == "__main__":
    main()
